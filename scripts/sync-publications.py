"""Sync manuscript figures, legends and complete tables without fitting models."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import argparse
import hashlib
import json
import re
import shutil
import subprocess
from lxml import etree as ET

import openpyxl
import pypdfium2 as pdfium

ROOT = Path(__file__).resolve().parents[1]
SITE = ROOT / 'site/research'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
RUNTIME = Path('/Users/chenmo/.cache/codex-runtimes/codex-primary-runtime/dependencies')
FIGURES = [
    ('Figure 1.png', '10_研究流程', 'figure-1', 'Main figures'),
    ('Figure 2.pdf', '01_关联图谱', 'figure-2', 'Main figures'),
    ('Figure 3.pdf', '03_关联矩阵', 'figure-3', 'Main figures'),
    ('Figure 4.pdf', '02_方差解释', 'figure-4', 'Main figures'),
    ('Figure 5.pdf', '04_暴露网络', 'figure-5', 'Main figures'),
    ('Figure Extended 1.pdf', '12_图谱流程', 'extended-1', 'Extended Data figures'),
    ('Figure Extended 2.pdf', '09_关联谱', 'extended-2', 'Extended Data figures'),
    ('Figure Extended 3.pdf', '05_年度重复', 'extended-3', 'Extended Data figures'),
    ('Figure Extended 4.pdf', '08_年龄交互', 'extended-4', 'Extended Data figures'),
    ('Figure Extended 5.pdf', '11_效应振动', 'extended-5', 'Extended Data figures'),
    ('Figure Extended 6.pdf', '06_调整模型', 'extended-6', 'Extended Data figures'),
    ('Figure Extended 7.pdf', '13_评分对照', 'extended-7', 'Extended Data figures'),
    ('Figure Extended 8.pdf', '07_测量一致', 'extended-8', 'Extended Data figures'),
    ('Figure Extended 9.pdf', '14_肺功能网', 'extended-9', 'Extended Data figures'),
    *[(f'Figure Supplementary {i}.pdf', name, f'supplementary-{i}', 'Supplementary figures')
      for i, name in enumerate(['15_样本构成', '16_显著分布', '17_糖化谱',
                                '18_年度异质', '19_系数调整', '20_环境关联'], 1)],
]
SECTION_NAMES = {
    'Table S8': ['Air and urinary measurement concordance', 'Environmental sensitivity analyses', 'Additional exposure modules'],
    'Table S9': ['Coefficient robustness', 'Adjustment-factor contributions'],
    'Table S10': ['Cross-sectional replication', 'Longitudinal follow-up', 'Incident disease'],
    'Table S13': ['Agreement between dietary recalls', 'Paired-day phenotype associations'],
    'Table S15': ['Exposure-pair relationships', 'Gain comparisons'],
    'Table S16': ['Exposure associations', 'Shared aging and memory associations',
                  'Memory models with aging-score adjustment', 'Association-profile comparisons',
                  'Discovery selection', 'Temporal validation'],
}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def paragraphs(path):
    with ZipFile(path) as archive:
        document = ET.fromstring(archive.read('word/document.xml'))
    return [
        ''.join(n.text or '' for n in p.findall('.//w:t', NS))
        for p in document.findall('.//w:body/w:p', NS)]


def read_legends(path):
    result = {}
    current = None
    pattern = r'^(Figure \d+|Extended Data Figure \d+|Supplementary Figure \d+|Table 1|Supplementary Table \d+)\. (.+)$'
    for text in paragraphs(path):
        match = re.match(pattern, text)
        if match:
            current = match[1]
            result[current] = {'label': current, 'title': match[2].rstrip('.'), 'legend': []}
        elif text in ['Extended Data figures', 'Supplementary figures', 'Table legends', 'Figure 5 node key']:
            current = None
        elif current and text.strip():
            result[current]['legend'].append(text)
    assert len(result) == 37, len(result)
    return result


def split_input(archive, selected, destination):
    """Select one worksheet for import, preserving its original XML and styles."""
    tree = ET.fromstring(archive.read('xl/workbook.xml'))
    sheets = tree.find('s:sheets', NS)
    chosen = list(sheets)[selected]
    relationship_id = chosen.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    for sheet in list(sheets):
        if sheet is not chosen:
            sheets.remove(sheet)
    for view in tree.findall('s:bookViews/s:workbookView', NS):
        view.set('activeTab', '0')
        view.set('firstSheet', '0')
    names = tree.find('s:definedNames', NS)
    if names is not None:
        for name in list(names):
            local = name.get('localSheetId')
            if local is not None and int(local) != selected:
                names.remove(name)
            elif local is not None:
                name.set('localSheetId', '0')
    rels = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
    worksheet_path = None
    for rel in list(rels):
        if rel.get('Type', '').endswith('/worksheet'):
            if rel.get('Id') != relationship_id:
                rels.remove(rel)
            else:
                target = rel.get('Target')
                worksheet_path = target.lstrip('/') if target.startswith('/') else 'xl/' + target
    assert worksheet_path
    types = ET.fromstring(archive.read('[Content_Types].xml'))
    for part in list(types):
        path = part.get('PartName', '').lstrip('/')
        if path.startswith('xl/worksheets/') and path != worksheet_path:
            types.remove(part)
    replacements = {'xl/workbook.xml': ET.tostring(tree, encoding='utf-8', xml_declaration=True),
                    'xl/_rels/workbook.xml.rels': ET.tostring(rels, encoding='utf-8', xml_declaration=True),
                    '[Content_Types].xml': ET.tostring(types, encoding='utf-8', xml_declaration=True)}
    with ZipFile(destination, 'w', ZIP_DEFLATED) as target:
        for name in archive.namelist():
            if name.startswith('xl/worksheets/') and name != worksheet_path:
                continue
            target.writestr(name, replacements.get(name, archive.read(name)))


def display(value, code):
    if value is None:
        return ''
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return str(value)
    if code == 'General':
        return str(value)
    if code == '[=0]"< precision";[<0.001]0.00E+00;0.000':
        return '< precision' if value == 0 else f'{value:.2E}' if value < 0.001 else f'{value:.3f}'
    positive = code.split(';')[0]
    positive = re.sub(r'\[[^]]*\]', '', positive)
    match = re.fullmatch(r'(#,##)?0(?:\.(0+))?(%|E\+00)?', positive, re.I)
    if not match:
        raise ValueError(f'Unsupported numeric display format: {code}')
    digits = len(match[2] or '')
    suffix = (match[3] or '').upper()
    if suffix == 'E+00':
        return f'{value:.{digits}E}'
    scaled = value * 100 if suffix == '%' else value
    return format(scaled, f'{"," if match[1] else ""}.{digits}f') + ('%' if suffix == '%' else '')


def table_preview(sheet):
    blocks = []
    current = []
    for row in sheet.iter_rows():
        cells = [display(cell.value, cell.number_format) for cell in row]
        if any(cells):
            current.append(cells)
        elif current:
            blocks.append(current)
            current = []
    if current:
        blocks.append(current)
    names = SECTION_NAMES.get(sheet.title, [''])
    assert len(names) == len(blocks), (sheet.title, len(blocks))
    result = []
    for name, block in zip(names, blocks):
        width = max(i + 1 for row in block for i, value in enumerate(row) if value)
        assert all(block[0][i] for i in range(width)), (sheet.title, 'Missing column heading')
        result.append({'title': name, 'columns': block[0][:width],
                       'rows': [row[:width] for row in block[1:]]})
    return result


def compare_export(sheet, destination):
    output = openpyxl.load_workbook(destination, read_only=True, data_only=False)
    assert output.sheetnames == [sheet.title], (destination, output.sheetnames)
    exported = output[sheet.title]
    if exported.max_row is None:
        exported.calculate_dimension(force=True)
    count = 0
    for a, b in zip(sheet.iter_rows(), exported.iter_rows(max_row=sheet.max_row, max_col=sheet.max_column)):
        for old, new in zip(a, b):
            assert old.value == new.value or old.value in [None, ''] and new.value in [None, ''], (sheet.title, old.coordinate, old.value, new.value)
            if old.value is not None:
                assert display(old.value, old.number_format) == display(new.value, new.number_format), (sheet.title, old.coordinate, 'Display format changed')
                count += 1
    assert exported.max_row == sheet.max_row and exported.max_column == sheet.max_column
    output.close()
    return count


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path, help='Manuscript project directory')
    parser.add_argument('--only-table', help='Export one worksheet during validation')
    parser.add_argument('--skip-excel', action='store_true', help='Reuse already verified single-table downloads')
    args = parser.parse_args()
    source = args.source
    work = source / 'analysis/102_网页同步'
    work.mkdir(exist_ok=True)
    for name in ['figures', 'tables']:
        (SITE / name).mkdir(exist_ok=True)
    dependencies = work / 'node_modules'
    if not dependencies.exists():
        dependencies.symlink_to(RUNTIME / 'node/node_modules', target_is_directory=True)
    builder = work / '02_拆分表格.mjs'
    if not builder.exists():
        builder.symlink_to(ROOT / 'scripts/export-table.mjs')
    manuscript = source / 'word材料/03_中文稿件/01_中文稿_v3.docx'
    legends_path = source / 'word材料/04_补充材料/Figure_and_Table_Legends.docx'
    workbook_path = source / 'word材料/04_补充材料/Supplementary_Tables_KNHANES.xlsx'
    source_paths = [manuscript, legends_path, workbook_path] + [source / '投稿所需的图片' / item[0] for item in FIGURES]
    source_hashes = {str(p): digest(p) for p in source_paths}
    legends = read_legends(legends_path)
    body = '\n'.join(paragraphs(manuscript))
    assert all(re.search(f'(?<![补充展据])图{i}', body) for i in range(1, 6))
    assert all(f'扩展数据图{i}' in body for i in range(1, 10))
    assert all(f'补充图{i}' in body for i in range(1, 7))
    figure_items = []
    if not args.only_table:
        for filename, stem, identifier, group in FIGURES:
            original = source / '投稿所需的图片' / filename
            destination = SITE / 'figures' / (stem + original.suffix)
            shutil.copyfile(original, destination)
            preview = SITE / 'figures' / (stem + '.png')
            if original.suffix == '.pdf':
                with pdfium.PdfDocument(original) as document:
                    assert len(document) == 1
                    page = document[0]
                    bitmap = page.render(scale=min(2.5, 2400 / max(page.get_width(), page.get_height())))
                    image = bitmap.to_pil().convert('RGB').copy()
                    bitmap.close()
                    page.close()
                image.save(preview, optimize=True)
            label = filename.replace('Figure Extended', 'Extended Data Figure').replace('Figure Supplementary', 'Supplementary Figure').rsplit('.', 1)[0]
            figure_items.append({'id': identifier, 'group': group, 'kind': 'figure', **legends[label],
                                 'file': 'figures/' + destination.name, 'preview': 'figures/' + preview.name,
                                 'source_file': filename, 'sha256': digest(destination),
                                 'preview_sha256': digest(preview)})
            assert digest(original) == digest(destination)
            print('Figure:', label, flush=True)
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    assert workbook.sheetnames == ['Table 1'] + [f'Table S{i}' for i in range(1, 17)]
    table_items = []
    checks = []
    input_path = work / '01_表格输入.xlsx'
    with ZipFile(workbook_path) as archive:
        for index, sheet in enumerate(workbook):
            if args.only_table and sheet.title != args.only_table:
                continue
            output = SITE / 'tables' / (sheet.title + '.xlsx')
            if not args.skip_excel:
                split_input(archive, index, input_path)
                staged_output = work / '22_单表输出.xlsx'
                preview_range = f'A1:{openpyxl.utils.get_column_letter(min(sheet.max_column, 6))}{min(sheet.max_row, 8)}'
                subprocess.run([str(RUNTIME / 'node/bin/node'), '--max-old-space-size=8192',
                    '--preserve-symlinks-main', str(builder), str(input_path), str(staged_output),
                    str(work / f'{index+3:02d}_表格预览.png'), preview_range], check=True)
                cells = compare_export(sheet, staged_output)
                shutil.copyfile(staged_output, output)
            else:
                cells = compare_export(sheet, output)
            sections = table_preview(sheet)
            table_id = 'table-1' if index == 0 else f'table-s{index}'
            payload = SITE / 'tables' / (sheet.title + '.json')
            payload.write_text(json.dumps({'sections': sections}, ensure_ascii=False, separators=(',', ':')))
            label = 'Table 1' if index == 0 else f'Supplementary Table {index}'
            table_items.append({'id': table_id, 'group': 'Tables' if index == 0 else 'Supplementary tables',
                'kind': 'table', **legends[label], 'file': 'tables/' + output.name,
                'source_file': output.name, 'sha256': digest(output),
                'data': 'tables/' + payload.name, 'data_sha256': digest(payload),
                'rows': sum(len(section['rows']) for section in sections), 'sections': len(sections)})
            checks.append({'sheet': sheet.title, 'cells': cells, 'rows': sheet.max_row, 'columns': sheet.max_column,
                           'source_values_and_display_unchanged': True, 'sha256': digest(output)})
            print('Table:', sheet.title, 'verified cells:', cells, flush=True)
    workbook.close()
    if input_path.exists():
        input_path.unlink()
    assert all(digest(Path(path)) == value for path, value in source_hashes.items())
    (work / '20_表格核验.json').write_text(json.dumps(checks, ensure_ascii=False, indent=2))
    if args.only_table:
        return
    assert len(figure_items) == 20 and len(table_items) == 17
    full_workbook = SITE / 'tables' / workbook_path.name
    shutil.copyfile(workbook_path, full_workbook)
    payload = {'figures': figure_items, 'tables': table_items,
               'workbook': {'file': 'tables/' + workbook_path.name, 'sha256': digest(full_workbook)}}
    encoded = json.dumps(payload, ensure_ascii=False, separators=(',', ':'))
    (SITE / 'publications.json').write_text(encoded)
    # The Shiny selector and the early homepage use the same ordered catalog.
    old_format = [{**item, 'title': item['label'] + ' · ' + item['title']} for item in figure_items]
    (ROOT / 'src/data/figures.json').write_text(json.dumps(old_format, ensure_ascii=False, indent=2))
    (ROOT / 'src/data/tables.json').write_text(json.dumps(table_items, ensure_ascii=False, indent=2))
    (work / '21_来源核验.json').write_text(json.dumps(source_hashes, ensure_ascii=False, indent=2))
    (ROOT / 'checks/publication_sync.json').write_text(json.dumps({'figures': 20, 'tables': 17,
        'source_files_unchanged': True, 'workbook_sha256': digest(full_workbook),
        'table_checks': checks}, ensure_ascii=False, indent=2))
    print('Synced 20 figures and 17 tables. Source documents unchanged.', flush=True)


if __name__ == '__main__':
    main()
